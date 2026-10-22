#!/usr/bin/env python3
"""
Classical Urdu Library builder.

Finds public-domain classical Urdu books on the Internet Archive (archive.org),
downloads them, turns loose page images into one PDF per book, and prepares
translation-ready text.

Commands
  search     Build/refresh catalog.csv from archive.org (no downloads)
  download   Download every catalogued book not yet downloaded
  prepare    OCR downloaded books -> searchable PDF + page-by-page text (.md, .docx)
  merge      Merge a folder of page images (from any lawful source) into one PDF
  run        search + download + prepare in one go
  status     Print progress counts

Examples
  python urdu_library.py run
  python urdu_library.py search --year-cutoff 1930
  python urdu_library.py download --author "Mirza Ghalib" --limit 5
  python urdu_library.py merge ./my_scans/some_book --out some_book.pdf

Only freely downloadable items are taken: lending-library / print-disabled
items on archive.org are skipped automatically.
"""
import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
AUTHORS_FILE = HERE / "authors.json"
SEARCH_API = "https://archive.org/services/search/v1/scrape"
META_API = "https://archive.org/metadata/{}"
DOWNLOAD = "https://archive.org/download/{}/{}"
USER_AGENT = "ClassicalUrduLibrary/1.0 (personal research; polite crawler)"
RESTRICTED_COLLECTIONS = {"inlibrary", "printdisabled", "lendinglibrary"}
FIELDS = ["identifier", "title", "creator", "date", "year", "language", "collection"]
CATALOG_COLS = ["identifier", "title", "creator", "matched_author", "author_died", "year",
                "needs_review", "url", "status", "pdf", "text_md", "note"]
PAUSE = 1.0  # seconds between requests, be polite


# ---------------------------------------------------------------- helpers
def http_get(url, binary=False, retries=4):
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(req, timeout=60) as r:
                data = r.read()
                return data if binary else data.decode("utf-8", "replace")
        except Exception as e:  # network hiccup: back off and retry
            if attempt == retries - 1:
                raise
            wait = 5 * (attempt + 1)
            print(f"  retry in {wait}s ({e})")
            time.sleep(wait)


def download_file(url, dest: Path):
    """Stream to dest with resume-safe .part file. Skips if dest exists."""
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    part = dest.with_suffix(dest.suffix + ".part")
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=120) as r, open(part, "wb") as f:
        shutil.copyfileobj(r, f, length=1 << 20)
    part.rename(dest)
    return dest


def as_list(v):
    if v is None:
        return []
    return v if isinstance(v, list) else [v]


def year_of(item):
    for key in ("year", "date"):
        for v in as_list(item.get(key)):
            m = re.search(r"(1[5-9]\d\d|20\d\d)", str(v))
            if m:
                return int(m.group(1))
    return None


def safe_name(s, n=80):
    s = re.sub(r"[^\w\-. ]+", "_", s, flags=re.UNICODE).strip()
    return s[:n] or "book"


def load_authors():
    return json.loads(AUTHORS_FILE.read_text(encoding="utf-8"))


def read_catalog(path):
    if not path.exists():
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {r["identifier"]: r for r in csv.DictReader(f)}


def write_catalog(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CATALOG_COLS)
        w.writeheader()
        for r in sorted(rows.values(), key=lambda r: (r.get("matched_author", ""), r.get("title", ""))):
            w.writerow({c: r.get(c, "") for c in CATALOG_COLS})


def scrape(query, max_items):
    """Yield search hits using archive.org's cursor-based scrape API."""
    cursor, got = None, 0
    while got < max_items:
        params = {"q": query, "fields": ",".join(FIELDS), "count": "1000"}
        if cursor:
            params["cursor"] = cursor
        data = json.loads(http_get(SEARCH_API + "?" + urllib.parse.urlencode(params)))
        for item in data.get("items", []):
            yield item
            got += 1
            if got >= max_items:
                return
        cursor = data.get("cursor")
        if not cursor:
            return
        time.sleep(PAUSE)


# ---------------------------------------------------------------- search
URDU = "(language:urd OR language:urdu)"


def cmd_search(args):
    catalog_path = Path(args.dir) / "catalog.csv"
    catalog = read_catalog(catalog_path)
    authors = load_authors()
    if args.author:
        authors = [a for a in authors if args.author.lower() in a["name"].lower()]
    found = 0

    def add(item, author):
        nonlocal found
        ident = item["identifier"]
        cols = set(as_list(item.get("collection")))
        if cols & RESTRICTED_COLLECTIONS:
            return
        yr = year_of(item)
        died = author["died"] if author else None
        # Flag for a human look: unknown year, or a modern printing (after 1960),
        # which may add copyrighted notes/editing, or be by a modern namesake.
        review = "yes" if (yr is None or yr > 1960) else ""
        if ident in catalog:
            return
        catalog[ident] = {
            "identifier": ident,
            "title": " / ".join(map(str, as_list(item.get("title"))))[:300],
            "creator": " / ".join(map(str, as_list(item.get("creator"))))[:200],
            "matched_author": author["name"] if author else "",
            "author_died": died or "",
            "year": yr or "",
            "needs_review": review,
            "url": f"https://archive.org/details/{ident}",
            "status": "found",
        }
        found += 1

    for a in authors:
        names = [n for n in a["search"]] + [a["name_urdu"]]
        creator_q = " OR ".join(f'creator:"{n}"' for n in names)
        q = f"mediatype:texts AND {URDU} AND ({creator_q})"
        print(f"Searching: {a['name']}")
        try:
            for item in scrape(q, args.max_per_author):
                add(item, a)
        except Exception as e:
            print(f"  search failed for {a['name']}: {e}")
        time.sleep(PAUSE)

    if not args.author and args.year_cutoff:
        # Everything in Urdu printed up to the cutoff year, whoever wrote it.
        q = f"mediatype:texts AND {URDU} AND year:[1500 TO {args.year_cutoff}]"
        print(f"Searching: all Urdu books printed up to {args.year_cutoff}")
        try:
            for item in scrape(q, args.max_old):
                add(item, None)
        except Exception as e:
            print(f"  search failed: {e}")

    write_catalog(catalog_path, catalog)
    print(f"\n{found} new books added. Catalog now has {len(catalog)} books: {catalog_path}")


# ---------------------------------------------------------------- download
def pick_files(files):
    names = {f["name"]: f for f in files}
    pdfs = [f for f in files if f["name"].lower().endswith(".pdf")]
    # Prefer original/text PDF, avoid tiny derivative chunks
    pdfs.sort(key=lambda f: (f.get("format") != "Text PDF", -int(f.get("size", 0) or 0)))
    pdf = pdfs[0]["name"] if pdfs else None
    jp2 = next((f["name"] for f in files if f["name"].endswith("_jp2.zip")), None)
    txt = next((f["name"] for f in files if f["name"].endswith("_djvu.txt")), None)
    return pdf, jp2, txt, names


def images_to_pdf(image_paths, out_pdf: Path):
    """Merge page images (any order-sortable names) into a single PDF."""
    from PIL import Image  # pip install pillow
    pages = []
    for p in image_paths:
        im = Image.open(p)
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        pages.append(im)
    if not pages:
        raise ValueError("no images to merge")
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    pages[0].save(out_pdf, save_all=True, append_images=pages[1:], resolution=300)
    return out_pdf


def natural_key(p):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(p))]


IMG_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".jp2", ".webp", ".bmp", ".gif"}


def cmd_download(args):
    base = Path(args.dir)
    catalog_path = base / "catalog.csv"
    catalog = read_catalog(catalog_path)
    todo = [r for r in catalog.values()
            if r["status"] in ("found", "failed")
            and (not args.author or args.author.lower() in r["matched_author"].lower())
            and (args.include_review or r["needs_review"] != "yes")]
    if args.limit:
        todo = todo[: args.limit]
    print(f"{len(todo)} books to download")
    for i, r in enumerate(todo, 1):
        ident = r["identifier"]
        folder = base / "books" / safe_name(r["matched_author"] or "Other") / ident
        print(f"[{i}/{len(todo)}] {ident}  {r['title'][:60]}")
        try:
            meta = json.loads(http_get(META_API.format(ident)))
            md = meta.get("metadata", {})
            if str(md.get("access-restricted-item", "")).lower() == "true" or \
               set(as_list(md.get("collection"))) & RESTRICTED_COLLECTIONS:
                r["status"], r["note"] = "skipped", "restricted (lending library)"
                continue
            pdf, jp2, txt, _ = pick_files(meta.get("files", []))
            (folder).mkdir(parents=True, exist_ok=True)
            (folder / "metadata.json").write_text(json.dumps(md, ensure_ascii=False, indent=2), "utf-8")
            out_pdf = folder / f"{ident}.pdf"
            if pdf:
                download_file(DOWNLOAD.format(ident, urllib.parse.quote(pdf)), out_pdf)
            elif jp2:
                z = download_file(DOWNLOAD.format(ident, urllib.parse.quote(jp2)), folder / jp2)
                pages_dir = folder / "pages"
                with zipfile.ZipFile(z) as zf:
                    zf.extractall(pages_dir)
                imgs = sorted((p for p in pages_dir.rglob("*") if p.suffix.lower() in IMG_EXT), key=natural_key)
                images_to_pdf(imgs, out_pdf)  # book complete -> one PDF
                z.unlink(missing_ok=True)
            else:
                r["status"], r["note"] = "skipped", "no PDF or page images offered"
                continue
            if txt:  # archive.org's own OCR text, sometimes usable
                try:
                    download_file(DOWNLOAD.format(ident, urllib.parse.quote(txt)), folder / "archive_ocr.txt")
                except Exception:
                    pass
            r["status"], r["pdf"], r["note"] = "downloaded", str(out_pdf.relative_to(base)), ""
        except Exception as e:
            r["status"], r["note"] = "failed", str(e)[:200]
            print(f"  failed: {e}")
        finally:
            write_catalog(catalog_path, catalog)  # save progress after every book
            time.sleep(PAUSE)


# ---------------------------------------------------------------- prepare (OCR)
def have(cmd):
    return shutil.which(cmd) is not None


def write_docx(pages, out_path, title):
    try:
        from docx import Document  # pip install python-docx
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        return None
    doc = Document()
    doc.add_heading(title, level=1)
    for n, text in enumerate(pages, 1):
        doc.add_heading(f"Page {n}", level=2)
        for line in text.strip().splitlines():
            if not line.strip():
                continue
            p = doc.add_paragraph(line)
            ppr = p._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")  # right-to-left paragraph
            ppr.append(bidi)
            for run in p.runs:
                rpr = run._r.get_or_add_rPr()
                rpr.append(OxmlElement("w:rtl"))
        doc.add_paragraph("Translation:")
    doc.save(out_path)
    return out_path


def prepare_pdf(pdf: Path, lang="urd", force=False):
    folder = pdf.parent
    stem = pdf.stem
    searchable = folder / f"{stem}.searchable.pdf"
    sidecar = folder / f"{stem}.ocr.txt"
    if force or not sidecar.exists():
        if not have("ocrmypdf"):
            raise RuntimeError("ocrmypdf not installed (see README)")
        cmd = ["ocrmypdf", "-l", lang, "--force-ocr", "--sidecar", str(sidecar),
               "--output-type", "pdf", "--optimize", "1", str(pdf), str(searchable)]
        subprocess.run(cmd, check=True)
    text = sidecar.read_text(encoding="utf-8", errors="replace")
    pages = text.split("\f")
    md = folder / f"{stem}.md"
    with open(md, "w", encoding="utf-8") as f:
        f.write(f"# {stem}\n\n")
        for n, t in enumerate(pages, 1):
            f.write(f"## Page {n}\n\n<div dir=\"rtl\">\n\n{t.strip()}\n\n</div>\n\n**Translation:**\n\n\n")
    write_docx(pages, folder / f"{stem}.docx", stem)
    return md, searchable


def cmd_prepare(args):
    base = Path(args.dir)
    catalog_path = base / "catalog.csv"
    catalog = read_catalog(catalog_path)
    todo = [r for r in catalog.values() if r["status"] == "downloaded"
            and (not args.author or args.author.lower() in r["matched_author"].lower())]
    if args.limit:
        todo = todo[: args.limit]
    print(f"{len(todo)} books to prepare")
    for i, r in enumerate(todo, 1):
        pdf = base / r["pdf"]
        print(f"[{i}/{len(todo)}] {pdf.name}")
        try:
            md, _ = prepare_pdf(pdf, args.lang, args.force)
            r["status"], r["text_md"], r["note"] = "ready", str(md.relative_to(base)), ""
        except Exception as e:
            r["note"] = f"prepare failed: {str(e)[:180]}"
            print(f"  failed: {e}")
        write_catalog(catalog_path, catalog)


# ---------------------------------------------------------------- merge (local)
def cmd_merge(args):
    folder = Path(args.folder)
    imgs = sorted((p for p in folder.rglob("*") if p.suffix.lower() in IMG_EXT), key=natural_key)
    out = Path(args.out or (folder.name + ".pdf"))
    images_to_pdf(imgs, out)
    print(f"Merged {len(imgs)} pages -> {out}")
    if args.ocr:
        md, searchable = prepare_pdf(out, args.lang)
        print(f"Text: {md}\nSearchable PDF: {searchable}")


def cmd_status(args):
    catalog = read_catalog(Path(args.dir) / "catalog.csv")
    counts = {}
    for r in catalog.values():
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    review = sum(1 for r in catalog.values() if r["needs_review"] == "yes")
    print(f"Books in catalog: {len(catalog)}  (needs review: {review})")
    for k, v in sorted(counts.items()):
        print(f"  {k:11} {v}")


def cmd_run(args):
    cmd_search(args)
    cmd_download(args)
    if have("ocrmypdf"):
        cmd_prepare(args)
    else:
        print("\nSkipping text preparation: ocrmypdf is not installed (see README).")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", default="library", help="where catalog.csv and books/ live (default: ./library)")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p, search=False, dl=False, prep=False):
        p.add_argument("--author", help="only this author (part of the name is enough)")
        p.add_argument("--limit", type=int, default=0, help="at most this many books")
        if search:
            p.add_argument("--year-cutoff", type=int, default=1930,
                           help="also catalog ALL Urdu books printed up to this year (0 = off)")
            p.add_argument("--max-per-author", type=int, default=2000)
            p.add_argument("--max-old", type=int, default=50000)
        if dl:
            p.add_argument("--include-review", action="store_true",
                           help="also download books flagged needs_review")
        if prep:
            p.add_argument("--lang", default="urd", help="Tesseract language (urd, or urd+fas+ara)")
            p.add_argument("--force", action="store_true", help="redo OCR even if done")

    common(sub.add_parser("search"), search=True)
    common(sub.add_parser("download"), dl=True)
    common(sub.add_parser("prepare"), prep=True)
    common(sub.add_parser("run"), search=True, dl=True, prep=True)
    m = sub.add_parser("merge")
    m.add_argument("folder")
    m.add_argument("--out")
    m.add_argument("--ocr", action="store_true", help="also make text for translation")
    m.add_argument("--lang", default="urd")
    sub.add_parser("status")

    args = ap.parse_args()
    Path(args.dir).mkdir(parents=True, exist_ok=True)
    {"search": cmd_search, "download": cmd_download, "prepare": cmd_prepare,
     "merge": cmd_merge, "status": cmd_status, "run": cmd_run}[args.cmd](args)


if __name__ == "__main__":
    main()
